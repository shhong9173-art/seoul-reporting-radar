"""완성차 모델별·공장별 월간 판매: 현대차·기아 IR 엑셀(공개). kr_stats_watch가 부르고 결과는 kr-stats 저장소에 함께 담는다.

원천(2026-10-30 탐침):
- 현대차 IR '판매실적'(https://www.hyundai.com/worldwide/ko/company/ir/ir-resources/sales-results): 연도별 목록 POST
  /wsvc/ww/salesPerformance.list.do(lang, year) → 'Unit Sales by Model'(국내 판매·국내 공장 수출, 모델·세대별)과
  'Global Plant Sales'(해외 공장별 현지 판매·수출) 엑셀. 그 달 엑셀은 보도자료(매달 1일)보다 2~4주 늦게 올라온다.
- 기아 IR '월별 판매실적 현황': GET /api/investors/business-sales-results → '차종별판매실적'(국내 공장 내수·수출, 모델별),
  '해외공장판매실적'(공장별 출고) 엑셀. 갱신이 들쑥날쑥하다(10-30 기준 8월까지).

규칙:
- 같은 회사·같은 구분(국내 판매·수출·공장)·같은 모델(세대 코드까지) 안에서만 비교한다.
- 새 달 값이 처음 보이면 그 달 신호를 만든다. 직전 3개월 평균과 전년 같은 달을 함께 본다(8월 휴가·명절 같은 계절 변동은 전년 같은 달로 거른다).
- 끊김(직전 3개월 평균 300대 이상이 5% 이하로)은 같은 이름 다른 세대가 막 늘었으면 '세대 교체로 보임'으로 낮춘다.
- 엑셀 값은 회사 발표 잠정치다. 단종·생산 중단·물량 조정 같은 이유는 엑셀에 없다 — 회사 확인 전에는 모른다.
- 조회 실패는 0이 아니다. 직전 값을 그대로 두고 미확인으로 표시한다.
"""
from __future__ import annotations

import io
import re
import unicodedata

HMC_PAGE = "https://www.hyundai.com/worldwide/ko/company/ir/ir-resources/sales-results"
HMC_LIST = "https://www.hyundai.com/wsvc/ww/salesPerformance.list.do"
HMC_HOST = "https://www.hyundai.com"
KIA_PAGE = "https://worldwide.kia.com/ko/company/investor-relations/library/performance-and-plans"
KIA_LIST = "https://worldwide.kia.com/api/investors/business-sales-results"
KIA_FILES = "https://worldwide.kia.com/files/"
MAKERS = {"hmc": "현대차", "kia": "기아"}
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"}
PAGES = {"hmc": HMC_PAGE, "kia": KIA_PAGE}

KEEP_MONTHS = 25
STOP_BASE = 300        # 끊김: 직전 3개월 평균이 이 이상이던 계열만
STOP_RATIO = 0.05      # 그 평균의 5% 이하
DROP_BASE = 300        # 급감: 직전 3개월 평균 하한(타스만처럼 대수는 작아도 기사가 되는 모델이 있다)
SURGE_BASE = 500       # 급증: 직전 3개월 평균 하한
DROP_RATIO = 0.5       # 평균·전년 같은 달 모두의 절반 아래
DROP_RATIO_NO_YOY = 0.3
SURGE_RATIO = 2.0
SURGE_MIN = 1000
NEW_MIN = 20           # 첫 판매: 이 대수 이상일 때만(시험·관용 소량 제외)
PLANT_BASE = 1000      # 공장 합계 급감 하한
PLANT_DROP = 0.7       # 공장 합계: 평균·전년 같은 달의 70% 아래(30% 넘게 감소)

# 보도 대조 검색어(영문 모델명 → 국문 차명). 모르는 이름은 영문 그대로.
KO = {"avante": "아반떼", "elantra": "아반떼", "sonata": "쏘나타", "grandeur": "그랜저", "santa-fe": "싼타페", "santa fe": "싼타페", "tucson": "투싼",
      "palisade": "팰리세이드", "kona": "코나", "casper": "캐스퍼", "venue": "베뉴", "ioniq": "아이오닉", "nexo": "넥쏘", "staria": "스타리아",
      "porter": "포터", "creta": "크레타", "exter": "엑스터", "stargazer": "스타게이저", "verna": "베르나", "i10": "i10", "i20": "i20", "i30": "i30",
      "morning": "모닝", "picanto": "모닝", "ray": "레이", "k3": "K3", "k4": "K4", "k5": "K5", "k8": "K8", "k9": "K9", "sorento": "쏘렌토",
      "sportage": "스포티지", "carnival": "카니발", "seltos": "셀토스", "niro": "니로", "bongo": "봉고", "tasman": "타스만", "telluride": "텔루라이드",
      "sonet": "쏘넷", "syros": "시로스", "carens": "카렌스", "stonic": "스토닉", "pegas": "페가스", "mohave": "모하비", "pv5": "PV5",
      "ev3": "EV3", "ev4": "EV4", "ev5": "EV5", "ev6": "EV6", "ev9": "EV9", "ev2": "EV2", "gv60": "GV60", "gv70": "GV70", "gv80": "GV80",
      "g70": "G70", "g80": "G80", "g90": "G90", "hb20": "HB20"}


def keyword(label):
    w = re.split(r"[\s(]", label.strip().lower(), maxsplit=1)[0]
    return KO.get(w, KO.get(w.replace("-", " "), label.split(" (")[0]))


MONTH_RX = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\.?$|^(1[0-2]|[1-9])\s*월$", re.I)
MONTH_NO = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
AGGREGATE = re.compile(r"^(sub-?total|total|grand total|ordinary vehicle|special vehicle)$", re.I)

HMC_SCOPE = {"Domestic": "국내 판매", "Export": "국내 공장 수출"}
HMC_PLANT = {"HMI": "인도", "HMMA": "미국 앨라배마", "HMGMA": "미국 조지아(HMGMA)", "HMMC": "체코", "HMTR": "튀르키예", "HMB": "브라질",
             "BHMC": "중국(베이징현대)", "HMMI": "인도네시아", "KMX": "기아 멕시코 위탁", "Singapore": "싱가포르", "Vietnam": "베트남",
             "CKD": "반조립(CKD)"}
HMC_PLANT_SUB = {"Domestic": "현지 판매", "Export": "수출", "Korea": "한국행"}
KIA_SCOPE = {"Domestic": "국내 판매", "Export(excl CKD)": "국내 공장 수출(CKD 제외)"}
KIA_PLANT = {"U.S.": "미국 조지아", "China": "중국", "Slovakia": "슬로바키아", "Mexico": "멕시코", "India": "인도"}


def text(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(v))).strip()


def number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    t = text(v).replace(",", "")
    return int(t) if re.fullmatch(r"-?\d+", t) else None


def month_cols(row):
    cols = {}
    for i, v in enumerate(row):
        t = text(v)
        m = MONTH_RX.match(t)
        if m:
            cols[i] = MONTH_NO[m.group(1).lower()] if m.group(1) else int(m.group(2))
    return cols if len(cols) >= 6 else None


def values(row, cols):
    out = {}
    for i, m in cols.items():
        n = number(row[i]) if i < len(row) else None
        if n is not None:
            out[m] = n
    return out


def cell(row, i):
    return text(row[i]) if i < len(row) else ""


def family(label):
    """세대 코드 앞 이름: 'Avante (CN8 HEV)' → 'avante hev', 'Kona (SX2 EV)' → 'kona ev' (같은 계열·같은 동력끼리 세대 교체를 본다)."""
    base = re.sub(r"\(.*?\)", " ", label).lower()
    power = re.search(r"\b(hev|phev|ev|n)\b", label.lower().replace("(", " ").replace(")", " "))
    return re.sub(r"\s+", " ", f"{base} {power.group(1) if power else ''}").strip()


def parse_hmc(rows, kind):
    """현대차 IR 엑셀(모델별·해외 공장별 같은 모양, 10-30 탐침): A열 비움, B열은 구획 머리(값 없음: Domestic·Export 또는 공장 코드)이거나
    하위 구분(값 있음: PC·RV·CV 또는 Domestic·Export·Korea)이거나 합계(Total·Grand Total), C열 모델·Sub-total, D열부터 1~12월."""
    cols, group, sub, out = None, "", "", []
    for row in rows:
        row = list(row)
        mc = month_cols(row)
        if mc:
            cols = mc
            continue
        if not cols:
            continue
        b, c = cell(row, 1), cell(row, 2)
        vals = values(row, cols)
        if b and not c and not vals:
            group, sub = b, ""
            continue
        if b and not c:
            if b.lower() == "total" and kind == "plant" and group:
                out.append({"group": group, "sub": "", "label": "공장 합계", "total": True, "values": vals})
            elif kind == "plant" and group and not AGGREGATE.match(b):  # 모델 이름 없이 구분만 있는 공장(HTBC)
                out.append({"group": group, "sub": b, "label": "전체", "total": False, "values": vals})
            continue
        if b:
            sub = b
        if not c or AGGREGATE.match(c) or not group:
            continue
        out.append({"group": group, "sub": sub if kind == "plant" else "", "label": c, "total": False, "values": vals})
    return out


def parse_kia_model(rows):
    """기아 차종별: D열 모델, 아래쪽 'Ordinary/Special Vehicle' 소계가 앞 행들의 구분, B열 'Domestic'·'Export(excl CKD)' 합계가 그 위 전체의 구획."""
    cols, pending, block, out = None, [], [], []
    for row in rows:
        row = list(row)
        mc = month_cols(row)
        if mc:
            cols = mc
            continue
        if not cols:
            continue
        b, d = cell(row, 1), cell(row, 3)
        vals = values(row, cols)
        if d and vals is not None and not b:
            if re.match(r"^(ordinary|special) vehicle$", d, re.I):
                special = d.lower().startswith("special")
                for x in pending:
                    x["sub"] = "Special" if special else "Ordinary"
                block += pending
                pending = []
            elif not AGGREGATE.match(d):
                pending.append({"label": d, "values": vals})
            continue
        if b:
            if b in KIA_SCOPE:
                for x in block + pending:
                    out.append({"group": b, "sub": x.get("sub", ""), "label": x["label"] + (" (특수)" if x.get("sub") == "Special" else ""),
                                "total": False, "values": x["values"]})
            block, pending = [], []
    return out


def parse_kia_plant(rows):
    """기아 해외공장(10-30 탐침): B열 'X Plant Sales (Ex-factory)' 머리, 구획마다 'CY 2026 | Total | Jan…' 머리행, C열 모델, B열 'X Plant(s)' 합계
    (합계 행은 아직 안 나온 달을 0으로 채운다 — last_month가 모델 행으로 마지막 달을 정한다)."""
    cols, group, out = None, "", []
    for row in rows:
        row = list(row)
        b, c = cell(row, 1), cell(row, 2)
        m = re.match(r"^(.+?)\s+plants?\s+sales", b, re.I)
        if m:
            group, cols = m.group(1).strip(), None
            continue
        mc = month_cols(row)
        if mc:
            cols = mc
            continue
        if not cols or not group:
            continue
        vals = values(row, cols)
        if c and not b:
            if not AGGREGATE.match(c):
                out.append({"group": group, "sub": "", "label": c, "total": False, "values": vals})
        elif b and re.match(r"^.+\s+plants?$", b, re.I):
            out.append({"group": group, "sub": "", "label": "공장 합계", "total": True, "values": vals})
    return out


def last_month(rows):
    """엑셀에 실린 마지막 달: 모델 행에 값이 있는 가장 늦은 달(합계 행의 미래 달 0은 세지 않는다)."""
    months = [m for x in rows if not x["total"] for m in x["values"]]
    return max(months) if months else None


def scope_of(maker, kind, x):
    g, s = x["group"], x["sub"]
    if maker == "hmc" and kind == "model":
        return HMC_SCOPE.get(g, g)
    if maker == "hmc":
        name = HMC_PLANT.get(g, g)
        return f"{name} 공장" + (f" {HMC_PLANT_SUB.get(s, s)}" if s else "")
    if kind == "model":
        return KIA_SCOPE.get(g, g)
    return f"{KIA_PLANT.get(g, g)} 공장 출고"


def series_id(maker, kind, x):
    return f"{maker}:{kind}:{x['group']}:{x['sub']}:{x['label']}"


def read_xlsx(content):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


PARSERS = {("hmc", "model"): lambda r: parse_hmc(r, "model"), ("hmc", "plant"): lambda r: parse_hmc(r, "plant"),
           ("kia", "model"): parse_kia_model, ("kia", "plant"): parse_kia_plant}


def ym_add(ym, k):
    y, m = int(ym[:4]), int(ym[5:7])
    t = y * 12 + (m - 1) + k
    return f"{t // 12}-{t % 12 + 1:02d}"


def fmt(n):
    return f"{n:,.0f}"


def already_low(mo, ym):
    """직전 달이 이미 급감·끊김 상태였나(그 전 3개월 평균의 절반 아래) — 줄어든 뒤 낮게 머무는 달마다 다시 띄우지 않는다."""
    p1 = mo.get(ym_add(ym, -1))
    pb = [mo.get(ym_add(ym, -i)) for i in (2, 3, 4)]
    return p1 is not None and None not in pb and sum(pb) > 0 and p1 < (sum(pb) / 3) * DROP_RATIO


def _base(mo, ym):
    prev = [mo.get(ym_add(ym, -i)) for i in (1, 2, 3)]
    return prev, (sum(prev) / 3 if None not in prev else None), mo.get(ym_add(ym, -12))


def _ratio(v, ref):
    return f"{v / ref:.0%}" if ref else "-"


def scope_totals(rows):
    """구분(국내 판매·수출·공장별 현지 판매 등)마다 모델 행을 더한 달별 합계."""
    tot = {}
    for s in rows:
        if s.get("total"):
            continue
        t = tot.setdefault(s["scope"], {})
        for k, v in s["months"].items():
            t[k] = t.get(k, 0) + (v or 0)
    return tot


def signals(series, maker, kind, ym):
    """그 달 신호. 같은 회사·구분 안에서만 비교한다.
    구분 전체가 함께 줄어든 달(파업·휴가·선적 일정)은 모델마다 띄우지 않고 '구분 전체 급감' 한 줄로 묶는다 — 구분 전체보다 확실히 더 줄어든 모델만 따로."""
    rows = [s for s in series if s["maker"] == maker and s["kind"] == kind]
    by_scope = {}
    for s in rows:
        by_scope.setdefault(s["scope"], []).append(s)
    totals = scope_totals(rows)
    scope_ratio, out = {}, []
    for scope, mo in totals.items():
        prev, base, yoy = _base(mo, ym)
        v = mo.get(ym) if ym in mo else None
        if v is None or not base:
            continue
        scope_ratio[scope] = v / base
        n = sum(1 for s in by_scope[scope] if not s.get("total") and (s["months"].get(ym) or 0) > 0)
        if kind == "model" and n >= 3 and base >= PLANT_BASE and v < base * PLANT_DROP and (yoy is None or v < yoy * PLANT_DROP) and not already_low(mo, ym):
            out.append({"id": f"cs:{maker}:{kind}:{scope}:{ym}:구분 전체 급감", "series": f"{maker}:{kind}:{scope}", "maker": MAKERS[maker],
                        "signal": "구분 전체 급감", "tier": "B", "period": ym, "scope": scope, "label": "전체", "value": v, "base": round(base), "yoy": yoy,
                        "text": f"{MAKERS[maker]} {scope} 전체(모델 {n}개 합) {ym} {fmt(v)}대 — 직전 3개월 평균 {fmt(base)}대의 {_ratio(v, base)}"
                                + (f", 전년 같은 달 {fmt(yoy)}대({_ratio(v, yoy)})" if yoy else ", 전년 같은 달 값 없음")
                                + ". 모델 대부분이 함께 줄었다(파업·휴가·선적 일정 여부는 회사 확인)"})
    for s in rows:
        mo = s["months"]
        v = mo.get(ym)
        if v is None:
            continue
        prev, base, yoy = _base(mo, ym)
        name = f"{MAKERS[maker]} {s['scope']} · {s['label']} {ym} {fmt(v)}대"
        ytxt = f", 전년 같은 달 {fmt(yoy)}대" if yoy is not None else ", 전년 같은 달 값 없음"
        sig = None
        fam = family(s["label"])
        sib = [o for o in by_scope.get(s["scope"], []) if o is not s and not o.get("total") and family(o["label"]) == fam
               and (o["months"].get(ym) or 0) >= max(NEW_MIN, (base or 0) * 0.1)
               and sum(o["months"].get(ym_add(ym, -i)) or 0 for i in (2, 3, 4)) <= (o["months"].get(ym) or 0) * 0.2]
        newer = max(sib, key=lambda o: o["months"].get(ym) or 0) if sib else None
        swap = (f"같은 이름 {newer['label']}이 {fmt(newer['months'][ym])}대로 늘어남 — 세대 교체로 보임" if newer else "")
        low_before = already_low(mo, ym)
        whole = scope_ratio.get(s["scope"])
        with_whole = whole is not None and whole < 0.8 and base and v / base >= whole * 0.7  # 구분 전체와 같이 줄었다
        export = "수출" in s["scope"]
        if s.get("total"):
            if low_before:
                pass
            elif base and base >= PLANT_BASE and v <= base * STOP_RATIO and prev[0] >= base * 0.2:
                sig = ("공장 출고 멈춤", "A", f"{name} — 직전 3개월 평균 {fmt(base)}대{ytxt}. 공장 전체 판매·출고가 거의 멈춤")
            elif base and base >= PLANT_BASE and v < base * PLANT_DROP and (yoy is None or v < yoy * PLANT_DROP):
                sig = ("공장 급감", "B", f"{name} — 직전 3개월 평균 {fmt(base)}대의 {_ratio(v, base)}{ytxt}" + (f"({_ratio(v, yoy)})" if yoy else ""))
        elif low_before and v < (base or 0):
            pass
        elif base is not None and base >= STOP_BASE and v <= base * STOP_RATIO and prev[0] >= base * 0.2:
            if newer:
                sig = ("세대 교체로 보임", "B", f"{name} — 직전 3개월 평균 {fmt(base)}대에서 끊김. {swap}")
            else:
                big = base >= PLANT_BASE and not export  # 수출은 선적 일정으로 한 달 비는 일이 잦다
                sig = ("판매 끊김", "A" if big else "B", f"{name} — 직전 3개월 평균 {fmt(base)}대{ytxt}. 판매(출고)가 거의 끊김 — 단종·생산 중단·물량 조정 여부는 회사 확인")
        elif base is not None and base >= DROP_BASE and v < base * DROP_RATIO and not (with_whole and not newer) and (
                (yoy is not None and v < yoy * DROP_RATIO) or (yoy is None and v < base * DROP_RATIO_NO_YOY)):
            detail = f"{name} — 직전 3개월 평균 {fmt(base)}대의 {_ratio(v, base)}{ytxt}" + (f"({_ratio(v, yoy)})" if yoy else "")
            if whole is not None and whole < 0.8:
                detail += f" · 같은 구분 전체는 {whole:.0%}"
            sig = ("세대 교체로 보임", "B", f"{detail}. {swap}") if newer else ("급감", "B", detail)
        elif (base is not None and base >= SURGE_BASE and min(prev) > 0 and v >= SURGE_MIN and v > base * SURGE_RATIO
              and (yoy is None or (yoy >= SURGE_BASE and v > yoy * SURGE_RATIO))):  # 막 나온 모델이 늘어나는 것(전년 소량)은 급증이 아니다
            sig = ("급증", "B", f"{name} — 직전 3개월 평균 {fmt(base)}대의 {v / base:.1f}배{ytxt}")
        elif v >= NEW_MIN and all((mo.get(k) or 0) == 0 for k in mo if k < ym) and any(k < ym for k in mo):
            sig = ("첫 판매", "B", f"{name} — 이 구분에서 처음 판매(앞선 {sum(1 for k in mo if k < ym)}개월 0)")
        if sig:
            out.append({"id": f"cs:{s['id']}:{ym}:{sig[0]}", "series": s["id"], "maker": MAKERS[maker], "signal": sig[0], "tier": sig[1],
                        "period": ym, "scope": s["scope"], "label": s["label"], "value": v, "base": None if base is None else round(base),
                        "yoy": yoy, "text": sig[2]})
    order = {"A": 0, "B": 1}
    out.sort(key=lambda x: (order[x["tier"]], -(x["base"] or x["value"] or 0)))
    return out


def sources(now, http, errors):
    """이번 해·지난해 엑셀 주소 목록. 목록 조회 실패는 errors에 남긴다(0건 아님)."""
    found = []
    for year in (now.year, now.year - 1):
        r, err = http.call("POST", HMC_LIST, data={"lang": "ko", "year": str(year)})
        if err:
            errors.append(f"현대차 목록 {year}: {err}")
            continue
        try:
            lst = (r.json().get("data") or {}).get("list") or []
        except ValueError:
            errors.append(f"현대차 목록 {year}: JSON 아님")
            continue
        for x in lst[:1]:
            for attr, kind in (("attrSalesModelValue", "model"), ("attrGlobalPlantValue", "plant")):
                if x.get(attr):
                    found.append({"maker": "hmc", "kind": kind, "year": year, "url": HMC_HOST + x[attr]})
    r, err = http.call("GET", KIA_LIST, params={"year": 0, "page": 0, "language": "ko"})
    if err:
        errors.append(f"기아 목록: {err}")
        return found
    try:
        lst = (r.json().get("data") or {}).get("list") or []
    except ValueError:
        errors.append("기아 목록: JSON 아님")
        return found
    for year in (now.year, now.year - 1):
        entries = [x for x in lst if x.get("type") == "sales" and x.get("year") == year]
        for x in entries[:1]:
            for f in x.get("files") or []:
                title = text(f.get("title"))
                kind = "model" if "차종별" in title else "plant" if "해외공장" in title else None
                if kind and f.get("path"):
                    found.append({"maker": "kia", "kind": kind, "year": year, "url": KIA_FILES + f["path"]})
    return found


def collect(prev, now, http, *, baseline):
    """(블록, 새 이벤트, 점검). prev는 이전 kr-stats 문서의 carSales 블록."""
    prev = prev if isinstance(prev, dict) else {}
    stamp = now.isoformat(timespec="seconds")
    files = dict(prev.get("files") or {})
    series = {s["id"]: {**s, "months": dict(s.get("months") or {})} for s in prev.get("series") or [] if isinstance(s, dict) and s.get("id")}
    errors, ok, fails, same = [], 0, 0, 0
    for src in sources(now, http, errors):
        url, maker, kind, year = src["url"], src["maker"], src["kind"], src["year"]
        old = files.get(url)
        ims = {"If-Modified-Since": old["lastModified"]} if old and old.get("lastModified") else {}
        r, err = http.call("GET", url, timeout=90, ok=(200, 304), headers={**UA, **ims})
        if err:
            fails += 1
            errors.append(f"{MAKERS[maker]} {year} {kind}: {err}")
            continue
        if r.status_code == 304:
            same += 1
            ok += 1
            continue
        try:
            rows = PARSERS[(maker, kind)](read_xlsx(r.content))
            last = last_month(rows)
            if not rows or not last:
                raise ValueError("행 없음")
        except Exception as e:  # noqa: BLE001 — 형식 변경은 미확인
            fails += 1
            errors.append(f"{MAKERS[maker]} {year} {kind}: 형식 {type(e).__name__} {str(e)[:40]}")
            continue
        ok += 1
        prefix = f"{year}-"
        for s in series.values():  # 이 파일이 다시 실은 해의 값은 새 파일로 바꾼다
            if s["maker"] == maker and s["kind"] == kind:
                s["months"] = {k: v for k, v in s["months"].items() if not k.startswith(prefix)}
        for x in rows:
            sid = series_id(maker, kind, x)
            s = series.setdefault(sid, {"id": sid, "maker": maker, "kind": kind, "scope": scope_of(maker, kind, x), "label": x["label"],
                                        "total": x["total"], "months": {}})
            for m, v in x["values"].items():
                if m <= last:
                    s["months"][f"{year}-{m:02d}"] = v
        files[url] = {"maker": maker, "kind": kind, "year": year, "latest": f"{year}-{last:02d}", "fetchedAt": stamp,
                      "lastModified": r.headers.get("Last-Modified") or r.headers.get("last-modified")}
    # 오래된 달·빈 계열 정리
    for s in list(series.values()):
        months = sorted(s["months"])[-KEEP_MONTHS:]
        s["months"] = {k: s["months"][k] for k in months}
        if not any(s["months"].values()):
            del series[s["id"]]
    keep_years = {now.year, now.year - 1}
    files = {u: f for u, f in files.items() if f.get("year") in keep_years}
    latest = {}
    for f in files.values():
        key = f"{f['maker']}:{f['kind']}"
        latest[key] = max(latest.get(key, ""), f["latest"])
    rows = sorted(series.values(), key=lambda s: s["id"])
    sigs, events = [], []
    prev_latest = prev.get("latest") or {}
    for key, ym in sorted(latest.items()):
        maker, kind = key.split(":")
        cur = signals(rows, maker, kind, ym)
        sigs += cur
        if not baseline and prev_latest and ym > prev_latest.get(key, ""):
            for x in cur:
                events.append({"id": x["id"], "kind": "모델별 판매", "signal": x["signal"], "tier": x["tier"], "period": ym,
                               "company": x["maker"], "detectedAt": stamp, "text": f"{x['text']} (IR 엑셀)", "url": PAGES[maker],
                               "keyword": "공장" if x["label"] == "공장 합계" else keyword(x["label"])})
    status = "ok" if not fails and not errors else ("failed" if not ok else "partial")
    check = {"source": "완성차 판매(현대차·기아 IR)", "status": status, "ok": ok, "failed": fails,
             "latest": " · ".join(f"{MAKERS[k.split(':')[0]]} {'모델' if k.endswith('model') else '공장'} {v}" for k, v in sorted(latest.items())),
             **({"error": "; ".join(errors[:3])} if errors else {})}
    block = {"files": files, "latest": latest, "series": rows, "signals": sigs, "checkedAt": stamp, "unchangedFiles": same,
             "rules": ["현대차·기아 IR 엑셀(회사 발표 잠정치). 그 달 엑셀은 매달 1일 판매 보도자료보다 2~4주 늦게 올라온다 — 국내 모델별 수치는 이미 보도됐을 수 있다.",
                       "같은 회사·같은 구분(국내 판매·국내 공장 수출·해외 공장)·같은 모델(세대 코드까지) 안에서만 비교한다.",
                       "급감·급증은 직전 3개월 평균과 전년 같은 달을 함께 본다. 끊김은 같은 이름 다른 세대가 막 늘었으면 세대 교체로 낮춘다.",
                       "이유(단종·생산 중단·물량 조정·재고)는 엑셀에 없다. 회사 확인 전에는 모른다."]}
    return block, events, check
